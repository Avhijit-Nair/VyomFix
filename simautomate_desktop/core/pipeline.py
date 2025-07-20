import os
import shutil
import datetime
import importlib.util
import pandas as pd
from core.pdf_utils import extract_images_from_pdf
from core.ai_processing import get_gemini_code_for_images
from core.compliance import calculate_deviations, analyze_compliance, generate_compliance_mapping_report, aggregate_all_deviations
from core.google_docs import create_or_update_report_doc
import glob

def process_simulation_pdf(sim_pdf_path, baseline_folder, compliance_folder, output_dir, api_key):
    # Find the baseline PDF (assume first PDF in baseline folder)
    baseline_pdfs = [f for f in os.listdir(baseline_folder) if f.lower().endswith('.pdf')]
    if not baseline_pdfs:
        raise FileNotFoundError('No baseline PDF found in baseline folder.')
    baseline_pdf_path = os.path.join(baseline_folder, baseline_pdfs[0])

    # Extract images from both PDFs
    real_images = extract_images_from_pdf(baseline_pdf_path, output_folder='extracted_images', suffix='real')
    sim_images = extract_images_from_pdf(sim_pdf_path, output_folder='extracted_images', suffix='simulated')
    if not real_images or not sim_images:
        raise RuntimeError('Failed to extract images from PDFs.')

    # Save results in a versioned folder
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    report_folder = os.path.join(output_dir, f'report_{timestamp}')
    os.makedirs(report_folder, exist_ok=True)

    # Only process the first 3 chart pairs for demo speed
    max_charts = 3
    df_real_list = []
    df_sim_list = []
    for i, (image1, image2) in enumerate(zip(real_images, sim_images)):
        if i >= max_charts:
            break
        code = get_gemini_code_for_images(image1, image2, api_key)
        code_file = f'flight_data_extraction_{i+1}.py'
        with open(code_file, 'w', encoding='utf-8') as f:
            f.write(code)
        spec = importlib.util.spec_from_file_location(f'flight_data_extraction_{i+1}', code_file)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        df_real = module.df_real_flight
        df_sim = module.df_simulated_flight
        df_real_list.append(df_real)
        df_sim_list.append(df_sim)
        # Save each chart's data
        df_real.to_csv(os.path.join(report_folder, f'df_chart_{i+1}_real.csv'), index=False)
        df_sim.to_csv(os.path.join(report_folder, f'df_chart_{i+1}_simulated.csv'), index=False)

    # For backward compatibility, use the first chart for compliance
    df_real = df_real_list[0]
    df_sim = df_sim_list[0]
    df_results = calculate_deviations(df_real, df_sim)
    compliance_report = analyze_compliance(df_results)
    # Generate compliance mapping report
    # Aggregate deviations for all parameters from all charts
    agg_results = aggregate_all_deviations(report_folder)
    if not agg_results.empty:
        mapping_report, mapping_summary = generate_compliance_mapping_report(agg_results)
        import pandas as pd
        pd.DataFrame(mapping_report).to_csv(os.path.join(report_folder, 'compliance_mapping_report.csv'), index=False)
        with open(os.path.join(report_folder, 'compliance_mapping_report.txt'), 'w', encoding='utf-8') as f:
            f.write(mapping_summary + '\n')
            for row in mapping_report:
                f.write(str(row) + '\n')
        # Output in Word format
        from docx import Document
        doc = Document()
        doc.add_heading('Compliance Mapping Report', 0)
        doc.add_paragraph(mapping_summary)
        table = doc.add_table(rows=1, cols=len(mapping_report[0]) if mapping_report else 1)
        if mapping_report:
            hdr_cells = table.rows[0].cells
            for i, key in enumerate(mapping_report[0].keys()):
                hdr_cells[i].text = key
            for row in mapping_report:
                row_cells = table.add_row().cells
                for i, val in enumerate(row.values()):
                    row_cells[i].text = str(val)
        doc.save(os.path.join(report_folder, 'compliance_mapping_report.docx'))

    # Save dataframes
    df_real.to_csv(os.path.join(report_folder, 'df_real_flight.csv'), index=False)
    df_sim.to_csv(os.path.join(report_folder, 'df_simulated_flight.csv'), index=False)
    df_results.to_csv(os.path.join(report_folder, 'df_results.csv'), index=False)
    # Save compliance report as JSON and TXT
    import json
    with open(os.path.join(report_folder, 'compliance_report.json'), 'w', encoding='utf-8') as f:
        json.dump(compliance_report, f, indent=2)
    with open(os.path.join(report_folder, 'compliance_report.txt'), 'w', encoding='utf-8') as f:
        f.write(str(compliance_report))
    # Optionally, copy the original PDFs
    shutil.copy2(sim_pdf_path, os.path.join(report_folder, os.path.basename(sim_pdf_path)))
    shutil.copy2(baseline_pdf_path, os.path.join(report_folder, os.path.basename(baseline_pdf_path)))

    # Create Google Doc for collaborative editing
    report_title = f"VyomFix Compliance Report {timestamp}"
    initial_content = str(compliance_report)
    table_csv_path = os.path.join(report_folder, 'compliance_mapping_report.csv')
    doc_id, doc_url = create_or_update_report_doc(report_title, initial_content, table_csv_path=table_csv_path)
    with open(os.path.join(report_folder, 'google_doc_url.txt'), 'w', encoding='utf-8') as f:
        f.write(doc_url)

    return report_folder, doc_url

def cleanup_generated_code_files():
    for f in glob.glob('flight_data_extraction_*.py'):
        try:
            os.remove(f)
        except Exception:
            pass 